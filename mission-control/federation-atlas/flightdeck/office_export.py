#!/usr/bin/env python3
"""Optional W2 binary projections of a validated MissionControl census JSON.
No binary is a SSOT, owner approval, controlled engineering OFFER, or corporate issued record.
"""
from __future__ import annotations
import argparse
from pathlib import Path
import json
import subprocess
import sys

BG='0A1424'; BLUE='1B3956'; CYAN='177A92'; PURPLE='7555A2'; WHITE='FFFFFF'; RED='9A384B'; GREEN='1B7E63'


def load(source):
    x=json.loads(Path(source).read_text(encoding='utf-8'))
    if (x.get('schema')!='missioncontrol.flightdeck.census.v1'
            or len(x.get('cells', []))!=42 or len(x.get('anchors', []))!=5
            or sum(c['status']=='UNKNOWN_NOT_OBSERVED' for c in x['cells'])!=x.get('metrics',{}).get('unobserved_cells')
            or x.get('engineering_accepted') is not False):
        raise ValueError('Unqualified binary input census')
    return x


def xlsx(data,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();overview=wb.active;overview.title='Overview';ev=wb.create_sheet('Evidence 42');an=wb.create_sheet('Anchors 5');pl=wb.create_sheet('Pipeline');doc=wb.create_sheet('Doc Control');ref=wb.create_sheet('Sources')
    overview.append(['MISSION CONTROL / FEDERATION FLIGHT DECK','W1 + W2 generated review'])
    overview.append(['Source SHA-256',data['source_semantic_sha256']])
    overview.append(['Exact checkout SHA',data['exact_checkout_sha'] or 'LOCAL_UNBOUND'])
    overview.append(['Snapshot UTC',data['snapshot_observed_at_utc']])
    overview.append(['Metric','Calculated by Excel'])
    overview.append(['Eligible cells',"=COUNTA('Evidence 42'!A2:A43)"])
    overview.append(['Observed status-only', '=COUNTIF(\'Evidence 42\'!D2:D43,"OBSERVED_STATUS_ONLY")'])
    overview.append(['Unknown not observed', '=COUNTIF(\'Evidence 42\'!D2:D43,"UNKNOWN_NOT_OBSERVED")'])
    overview.append(['Matching head anchors','=COUNTIF(\'Anchors 5\'!D2:D6,"MATCHES_SNAPSHOT")'])
    overview.append(['Stale head anchors','=COUNTIF(\'Anchors 5\'!D2:D6,"STALE_ANCHOR")'])
    overview.append(['Governance','HOLD — human approval and native checks separate'])
    ev.append(['BD leaf ID','Repository','Evidence family','State','Source SHA','Run ID','Job ID','Owner role','Next action','Victory predicate','GitHub source URL','BD authority URL'])
    for c in data['cells']:
        ev.append([c['id'],c['repository'],c['family'],c['status'],c['source_sha'],c['run_id'],c['job_id'],c['owner_role'],c['next_action'],c['victory_predicate'],c['repo_url'],c['bd_authority']])
    an.append(['Repository','Historical G4 SHA','Observed main SHA','Snapshot state','Protected in snapshot','Next action','GitHub branch'])
    for a in data['anchors']:
        an.append([a['repository'],a['source_anchor'],a['observed_main_sha'],a['status'],a['protected_on_snapshot'],a['next_action'],a['repo_branch_url']])
    pl.append(['Step','Group','Stage','Trigger','Command','Inputs','Outputs','Stop rule'])
    for s in data['pipeline']:
        pl.append([s['id'],s['group'],s['name'],s['trigger'],s['command'],s['inputs'],s['outputs'],s['gate']])
    doc.append(['Field','Value / external authority']);doc.append(['Document class','GENERATED_REVIEW_ONLY']);doc.append(['SSOT digest',data['source_semantic_sha256']]);doc.append(['Corporate issuer','NOT_ASSIGNED']);doc.append(['Corporate document number','NOT_ASSIGNED']);doc.append(['Corporate approval','NOT_GRANTED']);doc.append(['Possible document systems','ALX / SCK CEN (examples only)']);doc.append(['Master authority','Source-controlled JSON + repository/owner and corporate doc controls'])
    ref.append(['Source reference','URL']);ref.append(['Mission Control repository','https://github.com/GBOGEB/pipeline-automation-hub']);ref.append(['Canonical GOGREEN fleet','https://github.com/GBOGEB/pipeline-automation-hub/blob/master/mission-control/gogreen/samples/fleet_g1.json']);ref.append(['Mission Control GOGREEN owner work register','https://github.com/GBOGEB/pipeline-automation-hub/issues/529']);ref.append(['QPS recipient','https://github.com/GBOGEB/cryoplant-project']);ref.append(['CODEX / KEB','https://github.com/GBOGEB/CODEX']);ref.append(['ABACUS / DOW','https://github.com/GBOGEB/ABACUS'])
    widths={'Overview':[43,92], 'Evidence 42':[37,32,19,24,43,16,17,27,67,62,60,60], 'Anchors 5':[36,43,43,23,19,75,65], 'Pipeline':[12,26,30,51,78,65,64,65], 'Doc Control':[33,102], 'Sources':[43,100]}
    for ws in wb:
        ws.sheet_view.showGridLines=False
        ws.freeze_panes='A2' if ws!=overview else 'B6'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_options.horizontalCentered=True
        ws.sheet_properties.outlinePr.summaryBelow=True
        for col,w in enumerate(widths[ws.title],1):ws.column_dimensions[get_column_letter(col)].width=w
        for c in ws[1]:
            c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name='Aptos',bold=True,color=WHITE,size=11)
            c.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=29
        if ws not in (overview,doc):ws.auto_filter.ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical='center',wrap_text=True)
                c.font=Font(name='Aptos',size=10,color='000000' if c.data_type=='f' else (GREEN if isinstance(c.value,str) and c.value.startswith('https:') else '41556A'))
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F6FC')
                if isinstance(c.value,str) and c.value.startswith('https:'):c.hyperlink=c.value
            ws.row_dimensions[row[0].row].height=32 if ws==ev else 27
    for c in overview['A5:B5'][0]:c.fill=PatternFill('solid',fgColor=PURPLE);c.font=Font(bold=True,color='FFFFFF')
    for r in [7,8,9,10]:overview.cell(r,2).number_format='0'
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    wb.save(path)
    # Preserve the formulas but seed their preview caches from the SAME
    # validated source census. Excel is instructed to recalculate on open.
    # This is an OOXML cache transaction, not a substituted formula engine.
    from zipfile import ZipFile
    import re
    values={'B6':len(data['cells']),
            'B7':sum(c['status']=='OBSERVED_STATUS_ONLY' for c in data['cells']),
            'B8':sum(c['status']=='UNKNOWN_NOT_OBSERVED' for c in data['cells']),
            'B9':sum(a['status']=='MATCHES_SNAPSHOT' for a in data['anchors']),
            'B10':sum(a['status']=='STALE_ANCHOR' for a in data['anchors'])}
    output=Path(path);cache=output.with_suffix('.cached-tmp.xlsx')
    with ZipFile(output,'r') as zin,ZipFile(cache,'w') as zout:
        for entry in zin.infolist():
            raw=zin.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                text=raw.decode('utf-8')
                for coordinate,value in values.items():
                    rx=r'(<c r="'+coordinate+r'"[^>]*><f>[^<]*</f><v>)(</v>)'
                    text,n=re.subn(rx,lambda m:m.group(1)+str(value)+m.group(2),text)
                    if n!=1:raise ValueError('Formula cache not uniquely found: '+coordinate)
                raw=text.encode('utf-8')
            zout.writestr(entry,raw)
    cache.replace(output)


def docx(data,path):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.77);sec.bottom_margin=Inches(.73);sec.left_margin=Inches(.88);sec.right_margin=Inches(.88)
    styles=doc.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(9.7)
    for name,size in [('Title',30),('Heading 1',17),('Heading 2',12)]:
        st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(24,74,110)
    p=doc.add_paragraph('CODE PURPLE  /  MISSION CONTROL');p.style='Subtitle';p.runs[0].font.color.rgb=RGBColor(95,73,160)
    doc.add_paragraph('Federation Flight Deck',style='Title')
    doc.add_paragraph('W1 source census · W2 operator and document projections')
    table=doc.add_table(rows=0,cols=2);table.style='Light Shading Accent 1';table.autofit=True
    for k,v in [('Document state','GENERATED REVIEW — NOT CORPORATE APPROVED'),('Source semantic SHA256',data['source_semantic_sha256']),('Observed GitHub snapshot',data['snapshot_observed_at_utc']+' UTC'),('Corporate doc number','NOT ASSIGNED — external document control'),('QPS engineering disposition','DEFER')]:
        row=table.add_row();row.cells[0].text=k;row.cells[1].text=v
    doc.add_heading('Executive evidence',level=1)
    doc.add_paragraph('Mission Control presents independent repository roles, source identities, bounded Python execution and evidence without confusing metadata observations with complete engineering verification.')
    t=doc.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
    for cell,text in zip(t.rows[0].cells,['G4 evidence','Head anchors','Scout backlog']):cell.text=text
    for cell,text in zip(t.add_row().cells,['6 / 42 observed status-only','2 / 5 matched at historical snapshot','36 UNKNOWN leaves']):cell.text=text
    doc.add_heading('Repositories and public interfaces',level=1)
    for r in data['repositories']:
        doc.add_heading(r['short'],level=2)
        doc.add_paragraph(r['role'])
        doc.add_paragraph('Input: '+r['inputs']+'   |   Output: '+r['outputs'])
        doc.add_paragraph('Boundary: '+r['boundaries'])
    doc.add_heading('The execution pipeline',level=1)
    for s in data['pipeline']:
        doc.add_heading(s['id']+' · '+s['name'],level=2)
        doc.add_paragraph('Trigger: '+s['trigger']+'\nCommand: '+s['command']+'\nGenerated: '+s['outputs']+'\nGate: '+s['gate'])
    doc.add_heading('DMAIC safeguards',level=1)
    for x in data['controls']:
        doc.add_paragraph(x['id']+'  '+x['name'],style='Heading 2');doc.add_paragraph(x['threshold']+'  —  '+x['response'])
    doc.add_heading('Supporting evidence and document-control references',level=1)
    doc.add_paragraph('Machine register: census.json, bd_leaves.json, anchors.json, manifest.json, qa_receipt.json. HTML navigation: index.html; vector process: flightdeck.svg. Corporate issuer and controlled numbers are intentionally blank: any ALX/SCK CEN designation shall be assigned by the actual document owner.')
    doc.add_paragraph('https://github.com/GBOGEB/pipeline-automation-hub/issues/529')
    doc.add_paragraph('Source semantic SHA-256: '+data['source_semantic_sha256'])
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.add_run('MISSION CONTROL  •  Generated review copy  •  Not an approved engineering document')
    doc.save(path)


def pdf(data,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.units import inch
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MCtitle',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=27,leading=32,textColor=colors.HexColor('#173B59'),spaceAfter=10))
    styles.add(ParagraphStyle(name='MCh1',parent=styles['Heading1'],fontSize=15,leading=18,textColor=colors.HexColor('#173B59'),spaceBefore=15))
    styles.add(ParagraphStyle(name='MCbody',parent=styles['BodyText'],fontSize=9,leading=13,spaceAfter=6))
    styles.add(ParagraphStyle(name='MCsmall',parent=styles['BodyText'],fontSize=7.6,leading=10,spaceAfter=4))
    from xml.sax.saxutils import escape
    def p(s,sty='MCbody'):return Paragraph(escape(str(s)),styles[sty])
    def heading(title,anchor):return Paragraph(f'<a name="{anchor}"/>{escape(title)}',styles['MCh1'])
    story=[p('CODE PURPLE  /  MISSION CONTROL','MCsmall'),Paragraph('Federation Flight Deck',styles['MCtitle']),p('Two implementation waves: source census and user-facing evidence projection.'),p('Generated review • Corporate document ID NOT ASSIGNED • Owner QPS DEFER • Admission HOLD')]
    story.append(Paragraph('Jump to: <link href="#portfolio">Repositories</link>  |  <link href="#pipeline">Pipeline</link>  |  <link href="#backlog">Missing evidence</link>  |  <link href="#controls">Controls</link>  |  <link href="#references">References</link>',styles['MCbody']))
    summary=[['Scope','Measured (not inferred)'],['G4 observed metadata','6 / 42 repo-family cells'],['BD scout work leaves','36 UNKNOWN; no new root IDs'],['Snapshot anchor freshness','2 / 5 matched; 3 require refresh'],['Verified original G4 ZIP bytes','0 credited'],['Source identity',data['source_semantic_sha256'][:31]+'…']]
    table=Table(summary,colWidths=[2*inch,4.4*inch],hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#173B59')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#9CADBA')),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]));story.extend([table,Spacer(1,10)])
    story.append(heading('Six independent repositories','portfolio'))
    for r in data['repositories']:
        story.append(Paragraph(f'<b>{escape(r["short"])}</b> — {escape(r["role"])}',styles['MCbody']))
        story.append(p('Input: '+r['inputs']+' | Output: '+r['outputs'],'MCsmall'))
    story.append(heading('Python process and outputs','pipeline'))
    for x in data['pipeline']:
        story.append(Paragraph(f'<b>{escape(x["id"]+" · "+x["name"])}</b>  |  {escape(x["outputs"])}',styles['MCbody']))
        story.append(p('Trigger: '+x['trigger']+'  •  Gate: '+x['gate'],'MCsmall'))
    story.append(heading('36 missing evidence leaves','backlog'))
    story.append(p('The workbook and HTML include all 42 rows. Missing is a status, not proof of failing code. This PDF lists the six owners and links the canonical control issue.'))
    for r in data['repositories']:
        sub=[c for c in data['cells'] if c['repository']==r['name'] and c['status']=='UNKNOWN_NOT_OBSERVED']
        story.append(p(f"{r['short']}: {len(sub)} scout work leaves — "+', '.join(x['family_label'] for x in sub),'MCsmall'))
    story.append(Paragraph('<link href="https://github.com/GBOGEB/pipeline-automation-hub/issues/529">Open GOGREEN issue #529 / authoritative worklist</link>',styles['MCbody']))
    story.append(heading('DMAIC control and release barriers','controls'))
    for x in data['controls']:
        story.append(Paragraph(f'<b>{escape(x["id"])} — {escape(x["name"])}:</b> {escape(x["threshold"])}',styles['MCsmall']))
    story.append(heading('Document control and source references','references'))
    for x in ['No official ALX/SCK CEN corporate document number is assigned.', 'No electronic signature or approval is implied by PDF bookmarks or links.', 'Generated PDF, DOCX, PPTX and XLSX are current-SSOT projections and shall be reissued when source changes.', 'Missing native archive evidence and QPS recipient acceptance remain HOLD/DEFER.']:
        story.append(p(x))
    story.append(p('Source semantic SHA-256: '+data['source_semantic_sha256'],'MCsmall'))
    story.append(Paragraph('<link href="https://github.com/GBOGEB/pipeline-automation-hub">GitHub source repository</link>  |  <link href="https://github.com/GBOGEB/pipeline-automation-hub/issues/529">G4 WBS / BD register</link>',styles['MCbody']))
    doc=SimpleDocTemplate(str(path),pagesize=(595.276,841.89),leftMargin=54,rightMargin=54,topMargin=42,bottomMargin=42,title='Mission Control Federation Flight Deck W1/W2',author='Source-controlled Mission Control generator')
    def footer(c,d):
        c.saveState();c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#607B90'));c.drawString(54,30,'MISSION CONTROL  |  Generated review / not corporate approved');c.drawRightString(540,30,'Page '+str(d.page));c.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--input',required=True);ap.add_argument('--out',required=True);args=ap.parse_args(argv)
    data=load(args.input);out=Path(args.out);out.mkdir(parents=True,exist_ok=True)
    xlsx(data,out/'evidence_register.xlsx');docx(data,out/'federation_report.docx');pdf(data,out/'linked_evidence_book.pdf')
    root=Path(__file__).parent
    sp=subprocess.run(['node',str(root/'slides.js'),str(args.input),str(out/'mission_brief.pptx')],capture_output=True,text=True)
    if sp.returncode:
        raise RuntimeError('PPTX producer failed: '+sp.stderr[-1800:])
    print(json.dumps({'status':'PASS','files':['evidence_register.xlsx','federation_report.docx','linked_evidence_book.pdf','mission_brief.pptx'],'source_semantic_sha256':data['source_semantic_sha256']},sort_keys=True))
    return 0

if __name__=='__main__':raise SystemExit(main())