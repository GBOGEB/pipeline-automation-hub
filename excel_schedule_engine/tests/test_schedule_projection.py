import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from schedule_projection import ScheduleProjectionEngine


class ScheduleProjectionTests(unittest.TestCase):
    def write_candidate(self, root: Path, headers, rows, table_id="Plan__Schedule__abc123", sheet="Plan"):
        csv_dir = root / "Outputs" / "excel" / "tables_csv"
        csv_dir.mkdir(parents=True, exist_ok=True)
        csv_path = csv_dir / f"{table_id}.csv"
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        manifest_dir = root / "Outputs" / "excel"
        manifest_dir.mkdir(parents=True, exist_ok=True)
        manifest = {
            "schema": "gbogeb.excel_schedule_engine.table_manifest.v1",
            "authority": {"authority_transfer": False},
            "summary": {"tables_exported": 1, "schedule_candidates": 1, "errors": 0},
            "tables": [
                {
                    "id": table_id,
                    "source_sheet": sheet,
                    "schedule_candidate": True,
                    "csv_path": str(csv_path.relative_to(root)),
                }
            ],
        }
        manifest_path = manifest_dir / "table_manifest.json"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        return manifest_path

    def test_valid_schedule_maps_to_canonical_outputs_and_dashboard(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self.write_candidate(
                root,
                ["Task ID", "Task", "Start Date", "Finish Date", "Predecessors", "Owner", "Status", "% Complete", "Total Float", "Milestone"],
                [
                    ["A", "Design", "2026-09-22", "2026-09-24", "", "Alice", "Completed", "100%", "2", "No"],
                    ["B", "Build", "2026-09-25", "2026-09-28", "A", "Bob", "In Progress", "50%", "1", "No"],
                    ["M1", "Gate", "2026-09-29", "2026-09-29", "B", "PM", "Planned", "0%", "0", "Yes"],
                ],
            )
            engine = ScheduleProjectionEngine(root, as_of=date(2026, 9, 22))
            manifest_path = engine.run()

            payload = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(payload["summary"]["activities"], 3)
            self.assertEqual(payload["summary"]["red"], 0)
            self.assertGreaterEqual(payload["summary"]["green"], 1)
            self.assertFalse(payload["authority"]["authority_transfer"])

            canonical = list(csv.DictReader(engine.canonical_csv.open(encoding="utf-8-sig")))
            self.assertEqual(canonical[0]["activity_id"], "A")
            self.assertEqual(canonical[1]["predecessor_ids"], "A")
            self.assertEqual(canonical[2]["milestone"], "True")
            self.assertEqual(canonical[2]["duration_days"], "0.0")

            wb = load_workbook(engine.dashboard_xlsx)
            self.assertIn("Summary", wb.sheetnames)
            self.assertIn("Schedule", wb.sheetnames)
            self.assertIn("Findings", wb.sheetnames)
            wb.close()

    def test_duplicate_dates_dependencies_progress_and_milestone_fail_validation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self.write_candidate(
                root,
                ["ID", "Activity", "Start", "Finish", "Duration", "Predecessor", "Owner", "Status", "Progress", "Float", "Milestone"],
                [
                    ["A", "Bad dates", "2026-09-30", "2026-09-20", "-1", "Z", "", "Blocked", "120", "-2", "No"],
                    ["A", "Duplicate", "2026-09-22", "2026-09-23", "1", "", "Owner", "Open", "0", "1", "No"],
                    ["M", "Bad milestone", "2026-09-22", "2026-09-24", "2", "", "Owner", "Open", "0", "0", "Yes"],
                ],
            )
            engine = ScheduleProjectionEngine(root, as_of=date(2026, 9, 22))
            engine.run()
            findings = list(csv.DictReader(engine.findings_csv.open(encoding="utf-8-sig")))
            codes = {row["code"] for row in findings}
            for expected in {
                "DUPLICATE_ACTIVITY_ID",
                "START_AFTER_FINISH",
                "NEGATIVE_DURATION",
                "UNKNOWN_PREDECESSOR",
                "MISSING_OWNER",
                "PROGRESS_OUT_OF_RANGE",
                "MILESTONE_NONZERO_SPAN",
                "MILESTONE_NONZERO_DURATION",
            }:
                self.assertIn(expected, codes)
            kpis = json.loads(engine.kpis_json.read_text(encoding="utf-8"))
            self.assertGreater(kpis["ryg"]["red"], 0)
            self.assertGreater(kpis["validation"]["error_findings"], 0)

    def test_predecessor_cycle_is_detected(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self.write_candidate(
                root,
                ["Activity ID", "Activity", "Start Date", "Finish Date", "Predecessors", "Owner", "Status"],
                [
                    ["A", "One", "2026-09-22", "2026-09-23", "B", "A", "Open"],
                    ["B", "Two", "2026-09-23", "2026-09-24", "A", "B", "Open"],
                ],
            )
            engine = ScheduleProjectionEngine(root, as_of=date(2026, 9, 22))
            engine.run()
            findings = list(csv.DictReader(engine.findings_csv.open(encoding="utf-8-sig")))
            cycle_rows = [row for row in findings if row["code"] == "PREDECESSOR_CYCLE"]
            self.assertEqual(len(cycle_rows), 2)

    def test_missing_id_is_derived_and_overdue_is_red(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            self.write_candidate(
                root,
                ["Task", "Start Date", "Finish Date", "Owner", "Status"],
                [["Late task", "2026-09-01", "2026-09-10", "Owner", "In Progress"]],
            )
            engine = ScheduleProjectionEngine(root, as_of=date(2026, 9, 22))
            engine.run()
            canonical = list(csv.DictReader(engine.canonical_csv.open(encoding="utf-8-sig")))
            self.assertEqual(canonical[0]["id_origin"], "DERIVED")
            self.assertEqual(canonical[0]["ryg"], "RED")
            self.assertIn("OVERDUE", canonical[0]["risk_reasons"])
            findings = list(csv.DictReader(engine.findings_csv.open(encoding="utf-8-sig")))
            self.assertIn("MISSING_ACTIVITY_ID", {row["code"] for row in findings})
            self.assertIn("OVERDUE", {row["code"] for row in findings})

    def test_no_schedule_candidates_emits_empty_green_control_package(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            manifest_dir = root / "Outputs" / "excel"
            manifest_dir.mkdir(parents=True)
            manifest_path = manifest_dir / "table_manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "authority": {"authority_transfer": False},
                        "tables": [],
                        "summary": {"schedule_candidates": 0, "errors": 0},
                    }
                ),
                encoding="utf-8",
            )
            engine = ScheduleProjectionEngine(root, as_of=date(2026, 9, 22))
            output = engine.run()
            payload = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(payload["summary"]["activities"], 0)
            self.assertEqual(payload["summary"]["process_status"], "PASS")
            self.assertTrue(engine.dashboard_md.exists())
            self.assertTrue(engine.dashboard_xlsx.exists())


if __name__ == "__main__":
    unittest.main()
