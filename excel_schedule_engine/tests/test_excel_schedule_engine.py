import csv
import json
import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1] / "src"),
)
from excel_schedule_engine import ExcelScheduleEngine


class ExcelScheduleEngineTests(unittest.TestCase):
    def make_workbook(self, path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Plan"
        ws.append(
            [
                "Activity",
                "Start Date",
                "Finish Date",
                "Owner",
                "Status",
            ]
        )
        ws.append(
            [
                "Design",
                "2026-09-18",
                "2026-09-25",
                "GBO",
                "Open",
            ]
        )
        ws.append(
            [
                "Review",
                "2026-09-26",
                "2026-09-28",
                "Team",
                "Planned",
            ]
        )
        tab = Table(
            displayName="ScheduleTable",
            ref="A1:E3",
        )
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
        )
        ws.add_table(tab)

        ws2 = wb.create_sheet("Data")
        ws2.append(["Tag", "Value"])
        ws2.append(["A", 1])
        ws2.append(["B", 2])

        wb.create_sheet("Empty")
        wb.save(path)

    def test_exports_defined_table_and_used_range_separately(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            self.make_workbook(src)

            records = ExcelScheduleEngine(src, root).run()

            self.assertEqual(len(records), 2)
            plan = next(
                r for r in records if r.source_sheet == "Plan"
            )
            data = next(
                r for r in records if r.source_sheet == "Data"
            )

            self.assertEqual(
                plan.source_kind,
                "excel_table",
            )
            self.assertTrue(plan.schedule_candidate)
            self.assertGreaterEqual(
                plan.schedule_score,
                4,
            )
            self.assertEqual(
                data.source_kind,
                "sheet_used_range",
            )
            self.assertFalse(data.schedule_candidate)
            self.assertTrue(
                (root / plan.csv_path).exists()
            )
            self.assertTrue(
                (root / plan.xlsx_path).exists()
            )

            with (
                root / plan.csv_path
            ).open(
                encoding="utf-8-sig",
                newline="",
            ) as f:
                rows = list(csv.reader(f))

            self.assertEqual(
                rows[0][:3],
                [
                    "Activity",
                    "Start Date",
                    "Finish Date",
                ],
            )
            self.assertEqual(
                rows[1][0],
                "Design",
            )

            out_wb = load_workbook(
                root / plan.xlsx_path,
                data_only=False,
            )
            self.assertEqual(
                out_wb["Data"]["A2"].value,
                "Design",
            )
            self.assertEqual(
                out_wb["_META"]["B2"].value,
                "ScheduleTable",
            )
            self.assertIn(
                "ScheduleTable",
                out_wb["Data"].tables,
            )
            out_wb.close()

    def test_manifest_and_indexes_are_generated(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            self.make_workbook(src)

            ExcelScheduleEngine(
                src,
                root,
            ).run()

            manifest = json.loads(
                (
                    root
                    / "Outputs/excel/table_manifest.json"
                ).read_text(
                    encoding="utf-8"
                )
            )

            self.assertEqual(
                manifest["summary"]["tables_exported"],
                2,
            )
            self.assertEqual(
                manifest["summary"]["schedule_candidates"],
                1,
            )
            self.assertFalse(
                manifest["authority"]["authority_transfer"]
            )
            self.assertTrue(
                (
                    root
                    / "Reports/schedule_index.md"
                ).exists()
            )
            self.assertTrue(
                (
                    root
                    / "Reports/schedule_index.csv"
                ).exists()
            )

    def test_tables_only_skips_plain_used_range(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            self.make_workbook(src)

            records = ExcelScheduleEngine(
                src,
                root,
                include_sheet_ranges=False,
            ).run()

            self.assertEqual(
                len(records),
                1,
            )
            self.assertEqual(
                records[0].source_name,
                "ScheduleTable",
            )

    def test_fallback_starts_at_actual_used_range_origin(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Offset"
            ws["B3"] = "Task"
            ws["C3"] = "Start Date"
            ws["B4"] = "Build"
            ws["C4"] = "2026-09-18"
            wb.save(src)

            records = ExcelScheduleEngine(src, root).run()
            self.assertEqual(len(records), 1)
            record = records[0]
            self.assertEqual(record.source_ref, "B3:C4")
            self.assertEqual(
                record.headers,
                ["Task", "Start Date"],
            )
            self.assertEqual(record.data_rows, 1)
            self.assertTrue(record.schedule_candidate)

    def test_normalized_names_do_not_overwrite_each_other(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Plan A"
            ws1.append(["Task", "Status"])
            ws1.append(["A", "Open"])
            ws2 = wb.create_sheet("Plan_A")
            ws2.append(["Task", "Status"])
            ws2.append(["B", "Open"])
            wb.save(src)

            records = ExcelScheduleEngine(src, root).run()
            self.assertEqual(len(records), 2)
            self.assertEqual(
                len({r.id for r in records}),
                2,
            )
            self.assertEqual(
                len({r.csv_path for r in records}),
                2,
            )
            self.assertEqual(
                len({r.xlsx_path for r in records}),
                2,
            )
            for record in records:
                self.assertTrue(
                    (root / record.csv_path).exists()
                )
                self.assertTrue(
                    (root / record.xlsx_path).exists()
                )

    def test_formula_relocation_and_table_preservation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "MASTER.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Calc"
            ws["C5"] = "Item"
            ws["D5"] = "Qty"
            ws["E5"] = "Price"
            ws["F5"] = "Total"
            ws["C6"] = "Valve"
            ws["D6"] = 2
            ws["E6"] = 10
            ws["F6"] = "=D6*E6"
            tab = Table(
                displayName="CalcTable",
                ref="C5:F6",
            )
            tab.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showRowStripes=True,
            )
            ws.add_table(tab)
            wb.save(src)

            records = ExcelScheduleEngine(
                src,
                root,
                cell_mode="formula",
            ).run()
            record = records[0]
            out_wb = load_workbook(
                root / record.xlsx_path,
                data_only=False,
            )
            out_ws = out_wb["Data"]
            self.assertEqual(
                out_ws["D2"].value,
                "=B2*C2",
            )
            self.assertIn(
                "CalcTable",
                out_ws.tables,
            )
            out_wb.close()


if __name__ == "__main__":
    unittest.main()
