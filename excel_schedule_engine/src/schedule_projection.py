#!/usr/bin/env python3
"""Canonical planning/schedule projection, validation, KPI and RYG dashboard.

Consumes the governed table_manifest.json plus schedule-candidate CSV renditions.
The projection is a derived decision-support surface; it never mutates the source
workbook and never acquires source authority.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


AUTHORITY = "DERIVED_SCHEDULE_PROJECTION_NOT_SOURCE_AUTHORITY"

CANONICAL_FIELDS = [
    "source_table_id",
    "source_sheet",
    "source_row",
    "activity_key",
    "activity_id",
    "id_origin",
    "activity_name",
    "start_date",
    "finish_date",
    "duration_days",
    "duration_origin",
    "predecessor_ids",
    "owner",
    "status",
    "status_raw",
    "progress_pct",
    "total_float_days",
    "milestone",
    "ryg",
    "risk_reasons",
]

HEADER_ALIASES = {
    "activity_id": {
        "id", "activity id", "activity_id", "task id", "task_id", "uid",
        "unique id", "unique_id",
    },
    "activity_name": {
        "activity", "activity name", "activity_name", "task", "task name",
        "task_name", "work package", "work_package", "description",
    },
    "start_date": {
        "start", "start date", "start_date", "planned start", "baseline start",
    },
    "finish_date": {
        "finish", "finish date", "finish_date", "end", "end date", "end_date",
        "planned finish", "baseline finish",
    },
    "duration_days": {
        "duration", "duration days", "duration_days", "days", "working days",
        "lead time", "lead_time",
    },
    "predecessor_ids": {
        "predecessor", "predecessors", "dependency", "dependencies",
        "predecessor ids", "predecessor_ids",
    },
    "owner": {"owner", "responsible", "assignee", "lead", "resource"},
    "status": {"status", "state", "phase"},
    "progress_pct": {
        "progress", "% complete", "percent complete", "completion",
        "progress pct", "progress_pct",
    },
    "total_float_days": {
        "float", "slack", "total float", "total_float", "total float days",
        "total_float_days",
    },
    "milestone": {"milestone", "is milestone", "is_milestone", "milestone?"},
}

STATUS_ALIASES = {
    "NOT_STARTED": {"not started", "not_started", "planned", "plan", "open", "todo", "to do"},
    "IN_PROGRESS": {"in progress", "in_progress", "active", "started", "working"},
    "COMPLETE": {"complete", "completed", "done", "closed", "finished", "100%"},
    "ON_HOLD": {"on hold", "on_hold", "hold", "paused"},
    "CANCELLED": {"cancelled", "canceled", "cancel"},
    "BLOCKED": {"blocked", "blocked/issue", "issue"},
}

RYG_ORDER = {"GREEN": 0, "YELLOW": 1, "RED": 2}
RYG_FILLS = {
    "GREEN": "C6EFCE",
    "YELLOW": "FFEB9C",
    "RED": "FFC7CE",
}
SEVERITY_ORDER = {"INFO": 0, "WARN": 1, "ERROR": 2}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[_\-]+", " ", text.strip().casefold())
    return re.sub(r"\s+", " ", text)


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _map_headers(headers: Sequence[str]) -> dict[str, int]:
    normalized = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for field_name, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_norm_header(a) for a in aliases}
        for index, header in enumerate(normalized):
            if header in normalized_aliases:
                mapping[field_name] = index
                break
    return mapping


def _parse_date(value: object) -> date | None:
    text = _clean(value)
    if not text:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_float(value: object, percent: bool = False) -> float | None:
    text = _clean(value)
    if not text:
        return None
    is_percent = text.endswith("%")
    if is_percent:
        text = text[:-1].strip()
    text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    if percent and is_percent:
        return number
    return number


def _parse_bool(value: object) -> bool | None:
    text = _clean(value).casefold()
    if not text:
        return None
    if text in {"1", "true", "yes", "y", "x", "milestone"}:
        return True
    if text in {"0", "false", "no", "n"}:
        return False
    return None


def _normalize_status(value: object) -> tuple[str, str]:
    raw = _clean(value)
    folded = raw.casefold()
    if not raw:
        return "", raw
    for canonical, aliases in STATUS_ALIASES.items():
        if folded in aliases:
            return canonical, raw
    return "UNKNOWN", raw


def _split_predecessors(value: object) -> list[str]:
    raw = _clean(value)
    if not raw:
        return []
    tokens = [t.strip() for t in re.split(r"[;,|\n]+", raw) if t.strip()]
    result = []
    relation = re.compile(r"^(.+?):(?:FS|SS|FF|SF)(?:[+-]\s*\d+(?:\.\d+)?\s*[dD]?)?$", re.I)
    for token in tokens:
        match = relation.match(token)
        result.append((match.group(1) if match else token).strip())
    return result


def _iso(value: date | None) -> str:
    return value.isoformat() if value else ""


@dataclass
class Finding:
    source_table_id: str
    activity_key: str
    activity_id: str
    source_row: int
    kind: str
    severity: str
    code: str
    message: str


@dataclass
class ScheduleRow:
    source_table_id: str
    source_sheet: str
    source_row: int
    activity_key: str
    activity_id: str
    id_origin: str
    activity_name: str
    start_date: str
    finish_date: str
    duration_days: float | None
    duration_origin: str
    predecessor_ids: str
    owner: str
    status: str
    status_raw: str
    progress_pct: float | None
    total_float_days: float | None
    milestone: bool | None
    ryg: str = "GREEN"
    risk_reasons: str = ""
    _predecessor_list: list[str] = field(default_factory=list, repr=False)


class ScheduleProjectionEngine:
    def __init__(
        self,
        output_root: Path,
        table_manifest_path: Path | None = None,
        as_of: date | None = None,
        due_soon_days: int = 14,
    ):
        self.output_root = output_root
        self.table_manifest_path = (
            table_manifest_path
            or output_root / "Outputs" / "excel" / "table_manifest.json"
        )
        self.as_of = as_of or date.today()
        self.due_soon_days = due_soon_days
        self.schedule_dir = output_root / "Outputs" / "excel" / "schedule"
        self.reports_dir = output_root / "Reports"
        self.canonical_csv = self.schedule_dir / "canonical_schedule.csv"
        self.canonical_xlsx = self.schedule_dir / "canonical_schedule.xlsx"
        self.findings_csv = self.schedule_dir / "validation_findings.csv"
        self.kpis_json = self.schedule_dir / "schedule_kpis.json"
        self.schedule_manifest = self.schedule_dir / "schedule_manifest.json"
        self.dashboard_md = self.reports_dir / "schedule_dashboard.md"
        self.dashboard_xlsx = self.reports_dir / "schedule_dashboard.xlsx"

    def _prepare(self) -> None:
        self.schedule_dir.mkdir(parents=True, exist_ok=True)
        self.reports_dir.mkdir(parents=True, exist_ok=True)

    def _load_manifest(self) -> dict[str, Any]:
        payload = json.loads(self.table_manifest_path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("table manifest must be a JSON object")
        authority = payload.get("authority", {})
        if authority.get("authority_transfer") is not False:
            raise ValueError("table manifest authority guardrail is not false")
        return payload

    def _read_candidate_rows(
        self,
        table: dict[str, Any],
    ) -> tuple[list[str], list[list[str]]]:
        csv_path = self.output_root / str(table["csv_path"])
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            rows = list(reader)
        if not rows:
            return [], []
        return rows[0], rows[1:]

    def _finding(
        self,
        findings: list[Finding],
        row: ScheduleRow,
        kind: str,
        severity: str,
        code: str,
        message: str,
    ) -> None:
        findings.append(
            Finding(
                source_table_id=row.source_table_id,
                activity_key=row.activity_key,
                activity_id=row.activity_id,
                source_row=row.source_row,
                kind=kind,
                severity=severity,
                code=code,
                message=message,
            )
        )

    def _row_from_source(
        self,
        table: dict[str, Any],
        source_row: int,
        headers: Sequence[str],
        values: Sequence[str],
        mapping: dict[str, int],
        findings: list[Finding],
    ) -> ScheduleRow:
        def value(field_name: str) -> str:
            index = mapping.get(field_name)
            return values[index] if index is not None and index < len(values) else ""

        table_id = str(table["id"])
        activity_id = _clean(value("activity_id"))
        id_origin = "SOURCE"
        if not activity_id:
            activity_id = f"DERIVED-{source_row:05d}"
            id_origin = "DERIVED"
        activity_key = f"{table_id}:{activity_id}"

        start_raw = value("start_date")
        finish_raw = value("finish_date")
        start = _parse_date(start_raw)
        finish = _parse_date(finish_raw)
        duration_raw = value("duration_days")
        supplied_duration = _parse_float(duration_raw)
        duration = supplied_duration
        duration_origin = "SOURCE" if _clean(duration_raw) else ""
        if duration is None and start is not None and finish is not None and finish >= start:
            duration = float((finish - start).days)
            duration_origin = "DERIVED_CALENDAR_SPAN"

        status, status_raw = _normalize_status(value("status"))
        progress_raw = value("progress_pct")
        progress = _parse_float(progress_raw, percent=True)
        float_raw = value("total_float_days")
        total_float = _parse_float(float_raw)
        milestone_raw = value("milestone")
        milestone = _parse_bool(milestone_raw)
        predecessor_list = _split_predecessors(value("predecessor_ids"))

        row = ScheduleRow(
            source_table_id=table_id,
            source_sheet=str(table.get("source_sheet", "")),
            source_row=source_row,
            activity_key=activity_key,
            activity_id=activity_id,
            id_origin=id_origin,
            activity_name=_clean(value("activity_name")),
            start_date=_iso(start),
            finish_date=_iso(finish),
            duration_days=duration,
            duration_origin=duration_origin,
            predecessor_ids=";".join(predecessor_list),
            owner=_clean(value("owner")),
            status=status,
            status_raw=status_raw,
            progress_pct=progress,
            total_float_days=total_float,
            milestone=milestone,
            _predecessor_list=predecessor_list,
        )

        if id_origin == "DERIVED":
            self._finding(
                findings, row, "VALIDATION", "WARN", "MISSING_ACTIVITY_ID",
                "No source activity/task ID; deterministic row-local ID was derived.",
            )
        if not row.activity_name:
            self._finding(
                findings, row, "VALIDATION", "WARN", "MISSING_ACTIVITY_NAME",
                "Activity/task name is empty or no recognized activity-name column was mapped.",
            )
        if _clean(start_raw) and start is None:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "INVALID_START_DATE",
                f"Unparseable start date: {start_raw!r}.",
            )
        if _clean(finish_raw) and finish is None:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "INVALID_FINISH_DATE",
                f"Unparseable finish date: {finish_raw!r}.",
            )
        if start and finish and start > finish:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "START_AFTER_FINISH",
                "Start date is after finish date.",
            )
        if _clean(duration_raw) and supplied_duration is None:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "INVALID_DURATION",
                f"Unparseable duration: {duration_raw!r}.",
            )
        if supplied_duration is not None and supplied_duration < 0:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "NEGATIVE_DURATION",
                "Duration cannot be negative.",
            )
        if duration_origin == "DERIVED_CALENDAR_SPAN":
            self._finding(
                findings, row, "VALIDATION", "INFO", "DURATION_DERIVED_CALENDAR",
                "Duration was derived as calendar finish-start span; no working-calendar claim is made.",
            )
        if not row.owner:
            self._finding(
                findings, row, "VALIDATION", "WARN", "MISSING_OWNER",
                "Owner/responsible field is empty or unmapped.",
            )
        if not status:
            self._finding(
                findings, row, "VALIDATION", "WARN", "MISSING_STATUS",
                "Status/state field is empty or unmapped.",
            )
        elif status == "UNKNOWN":
            self._finding(
                findings, row, "VALIDATION", "WARN", "UNKNOWN_STATUS",
                f"Status value is not mapped to a canonical state: {status_raw!r}.",
            )
        if _clean(progress_raw) and progress is None:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "INVALID_PROGRESS",
                f"Unparseable progress: {progress_raw!r}.",
            )
        elif progress is not None and not 0 <= progress <= 100:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "PROGRESS_OUT_OF_RANGE",
                "Progress must be between 0 and 100 percent.",
            )
        if _clean(float_raw) and total_float is None:
            self._finding(
                findings, row, "VALIDATION", "ERROR", "INVALID_TOTAL_FLOAT",
                f"Unparseable total float: {float_raw!r}.",
            )
        if _clean(milestone_raw) and milestone is None:
            self._finding(
                findings, row, "VALIDATION", "WARN", "UNKNOWN_MILESTONE_FLAG",
                f"Milestone flag is not a recognized boolean: {milestone_raw!r}.",
            )
        if milestone is True:
            if start and finish and start != finish:
                self._finding(
                    findings, row, "VALIDATION", "ERROR", "MILESTONE_NONZERO_SPAN",
                    "Milestone has different start and finish dates.",
                )
            if duration is not None and duration > 0:
                self._finding(
                    findings, row, "VALIDATION", "ERROR", "MILESTONE_NONZERO_DURATION",
                    "Milestone duration is greater than zero.",
                )

        return row

    def _validate_ids_and_dependencies(
        self,
        rows: list[ScheduleRow],
        findings: list[Finding],
    ) -> None:
        by_table: dict[str, list[ScheduleRow]] = defaultdict(list)
        for row in rows:
            by_table[row.source_table_id].append(row)

        for table_id, table_rows in by_table.items():
            counts = Counter(
                row.activity_id for row in table_rows if row.id_origin == "SOURCE"
            )
            duplicate_ids = {activity_id for activity_id, count in counts.items() if count > 1}
            for row in table_rows:
                if row.activity_id in duplicate_ids and row.id_origin == "SOURCE":
                    self._finding(
                        findings, row, "VALIDATION", "ERROR", "DUPLICATE_ACTIVITY_ID",
                        f"Duplicate source activity ID within logical table: {row.activity_id!r}.",
                    )

            ids = {row.activity_id for row in table_rows}
            unique_rows = {
                row.activity_id: row
                for row in table_rows
                if counts.get(row.activity_id, 1) == 1
            }
            graph: dict[str, list[str]] = {}
            for row in table_rows:
                resolved: list[str] = []
                for predecessor in row._predecessor_list:
                    if predecessor == row.activity_id:
                        self._finding(
                            findings, row, "VALIDATION", "ERROR", "SELF_PREDECESSOR",
                            "Activity references itself as predecessor.",
                        )
                        continue
                    if predecessor not in ids:
                        self._finding(
                            findings, row, "VALIDATION", "ERROR", "UNKNOWN_PREDECESSOR",
                            f"Predecessor does not resolve within logical table: {predecessor!r}.",
                        )
                        continue
                    resolved.append(predecessor)
                if row.activity_id in unique_rows:
                    graph[row.activity_id] = resolved

            visiting: set[str] = set()
            visited: set[str] = set()
            cycle_nodes: set[str] = set()

            def visit(node: str, path: list[str]) -> None:
                if node in visiting:
                    if node in path:
                        cycle_nodes.update(path[path.index(node):])
                    return
                if node in visited:
                    return
                visiting.add(node)
                path.append(node)
                for predecessor in graph.get(node, []):
                    visit(predecessor, path)
                path.pop()
                visiting.remove(node)
                visited.add(node)

            for node in graph:
                visit(node, [])

            for node in cycle_nodes:
                row = unique_rows.get(node)
                if row:
                    self._finding(
                        findings, row, "VALIDATION", "ERROR", "PREDECESSOR_CYCLE",
                        "Activity participates in a predecessor cycle.",
                    )

    def _apply_risk_and_ryg(
        self,
        rows: list[ScheduleRow],
        findings: list[Finding],
    ) -> None:
        findings_by_key: dict[str, list[Finding]] = defaultdict(list)
        for finding in findings:
            findings_by_key[finding.activity_key].append(finding)

        due_soon_limit = self.as_of + timedelta(days=self.due_soon_days)
        for row in rows:
            reasons: list[str] = []
            red = any(
                f.severity == "ERROR" for f in findings_by_key.get(row.activity_key, [])
            )
            yellow = any(
                f.severity == "WARN" for f in findings_by_key.get(row.activity_key, [])
            )
            finish = _parse_date(row.finish_date)
            status = row.status
            active = status not in {"COMPLETE", "CANCELLED"}

            if active and finish and finish < self.as_of:
                red = True
                reasons.append("OVERDUE")
                self._finding(
                    findings, row, "RISK", "WARN", "OVERDUE",
                    f"Finish date {finish.isoformat()} is before as-of {self.as_of.isoformat()}.",
                )
            elif active and finish and self.as_of <= finish <= due_soon_limit:
                yellow = True
                reasons.append("DUE_SOON")
                self._finding(
                    findings, row, "RISK", "INFO", "DUE_SOON",
                    f"Finish date is within {self.due_soon_days} days of the as-of date.",
                )

            if row.total_float_days is not None:
                if row.total_float_days < 0:
                    red = True
                    reasons.append("NEGATIVE_FLOAT")
                    self._finding(
                        findings, row, "RISK", "WARN", "NEGATIVE_FLOAT",
                        "Total float is negative.",
                    )
                elif row.total_float_days == 0:
                    yellow = True
                    reasons.append("ZERO_FLOAT")

            if status == "BLOCKED":
                red = True
                reasons.append("BLOCKED")
            elif status == "ON_HOLD":
                yellow = True
                reasons.append("ON_HOLD")

            if red:
                row.ryg = "RED"
            elif yellow:
                row.ryg = "YELLOW"
            else:
                row.ryg = "GREEN"
            row.risk_reasons = ";".join(sorted(set(reasons)))

    def _calculate_kpis(
        self,
        rows: Sequence[ScheduleRow],
        findings: Sequence[Finding],
        candidate_tables: int,
    ) -> dict[str, Any]:
        total = len(rows)
        ryg = Counter(row.ryg for row in rows)
        statuses = Counter(row.status or "MISSING" for row in rows)
        error_keys = {
            finding.activity_key for finding in findings if finding.severity == "ERROR"
        }
        warning_keys = {
            finding.activity_key for finding in findings if finding.severity == "WARN"
        }
        progress_values = [
            row.progress_pct for row in rows if row.progress_pct is not None
        ]
        required_values = []
        for row in rows:
            required_values.extend(
                [
                    bool(row.activity_name),
                    bool(row.start_date),
                    bool(row.finish_date),
                    bool(row.owner),
                    bool(row.status),
                ]
            )
        validation_pass = total - len(error_keys)
        overdue = sum("OVERDUE" in row.risk_reasons.split(";") for row in rows)
        return {
            "schema": "gbogeb.schedule_kpis.v1",
            "as_of": self.as_of.isoformat(),
            "authority": AUTHORITY,
            "candidate_tables": candidate_tables,
            "activities_total": total,
            "ryg": {
                "green": ryg["GREEN"],
                "yellow": ryg["YELLOW"],
                "red": ryg["RED"],
            },
            "status_counts": dict(sorted(statuses.items())),
            "validation": {
                "activities_with_errors": len(error_keys),
                "activities_with_warnings": len(warning_keys),
                "findings_total": len(findings),
                "error_findings": sum(f.severity == "ERROR" for f in findings),
                "warning_findings": sum(f.severity == "WARN" for f in findings),
                "info_findings": sum(f.severity == "INFO" for f in findings),
                "validation_pass_rate_pct": round(
                    100.0 * validation_pass / total, 1
                ) if total else 100.0,
            },
            "planning": {
                "overdue": overdue,
                "milestones": sum(row.milestone is True for row in rows),
                "negative_float": sum(
                    row.total_float_days is not None and row.total_float_days < 0
                    for row in rows
                ),
                "zero_float": sum(row.total_float_days == 0 for row in rows),
                "mean_progress_pct": round(
                    sum(progress_values) / len(progress_values), 1
                ) if progress_values else None,
                "field_completeness_pct": round(
                    100.0 * sum(required_values) / len(required_values), 1
                ) if required_values else 100.0,
            },
        }

    def _write_csvs(
        self,
        rows: Sequence[ScheduleRow],
        findings: Sequence[Finding],
    ) -> None:
        with self.canonical_csv.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=CANONICAL_FIELDS)
            writer.writeheader()
            for row in rows:
                data = asdict(row)
                data.pop("_predecessor_list", None)
                writer.writerow({field_name: data.get(field_name) for field_name in CANONICAL_FIELDS})

        finding_fields = [
            "source_table_id", "activity_key", "activity_id", "source_row",
            "kind", "severity", "code", "message",
        ]
        with self.findings_csv.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=finding_fields)
            writer.writeheader()
            for finding in findings:
                writer.writerow(asdict(finding))

    @staticmethod
    def _autosize(ws) -> None:
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            width = max(
                10,
                min(
                    50,
                    max(
                        len(str(cell.value)) if cell.value is not None else 0
                        for cell in column
                    ) + 2,
                ),
            )
            ws.column_dimensions[letter].width = width

    def _write_canonical_xlsx(
        self,
        rows: Sequence[ScheduleRow],
        findings: Sequence[Finding],
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Schedule"
        ws.append(CANONICAL_FIELDS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            data = asdict(row)
            data.pop("_predecessor_list", None)
            ws.append([data.get(field_name) for field_name in CANONICAL_FIELDS])
            fill = PatternFill("solid", fgColor=RYG_FILLS[row.ryg])
            ws.cell(row=ws.max_row, column=CANONICAL_FIELDS.index("ryg") + 1).fill = fill
        ws.freeze_panes = "A2"
        self._autosize(ws)

        findings_ws = wb.create_sheet("Validation")
        fields = [
            "source_table_id", "activity_key", "activity_id", "source_row",
            "kind", "severity", "code", "message",
        ]
        findings_ws.append(fields)
        for cell in findings_ws[1]:
            cell.font = Font(bold=True)
        for finding in findings:
            findings_ws.append([getattr(finding, field_name) for field_name in fields])
        findings_ws.freeze_panes = "A2"
        self._autosize(findings_ws)

        meta = wb.create_sheet("_META")
        meta["A1"] = "authority"
        meta["B1"] = AUTHORITY
        meta["A2"] = "source_table_manifest"
        meta["B2"] = str(self.table_manifest_path)
        meta["A3"] = "as_of"
        meta["B3"] = self.as_of.isoformat()
        meta.sheet_state = "hidden"
        wb.save(self.canonical_xlsx)

    def _write_dashboard(
        self,
        rows: Sequence[ScheduleRow],
        findings: Sequence[Finding],
        kpis: dict[str, Any],
    ) -> None:
        lines = [
            "# Planning / Schedule Health Dashboard",
            "",
            f"As-of: **{self.as_of.isoformat()}**",
            "",
            f"Authority: `{AUTHORITY}`",
            "",
            "## Executive RYG",
            "",
            "| KPI | Value |",
            "|---|---:|",
            f"| Activities | {kpis['activities_total']} |",
            f"| GREEN | {kpis['ryg']['green']} |",
            f"| YELLOW | {kpis['ryg']['yellow']} |",
            f"| RED | {kpis['ryg']['red']} |",
            f"| Validation pass rate | {kpis['validation']['validation_pass_rate_pct']}% |",
            f"| Field completeness | {kpis['planning']['field_completeness_pct']}% |",
            f"| Overdue | {kpis['planning']['overdue']} |",
            f"| Negative float | {kpis['planning']['negative_float']} |",
            f"| Milestones | {kpis['planning']['milestones']} |",
            "",
            "## Attention list",
            "",
            "| RYG | Activity ID | Activity | Owner | Finish | Progress | Reasons |",
            "|---|---|---|---|---|---:|---|",
        ]
        attention = sorted(
            [row for row in rows if row.ryg != "GREEN"],
            key=lambda row: (-RYG_ORDER[row.ryg], row.finish_date or "9999-12-31", row.activity_key),
        )
        for row in attention[:50]:
            progress = "" if row.progress_pct is None else f"{row.progress_pct:g}%"
            lines.append(
                f"| {row.ryg} | {row.activity_id} | {row.activity_name or '-'} | "
                f"{row.owner or '-'} | {row.finish_date or '-'} | {progress or '-'} | "
                f"{row.risk_reasons or '-'} |"
            )
        if not attention:
            lines.append("| GREEN | - | No RED/YELLOW activities | - | - | - | - |")

        lines.extend(
            [
                "",
                "## Validation findings",
                "",
                "| Severity | Code | Activity | Message |",
                "|---|---|---|---|",
            ]
        )
        sorted_findings = sorted(
            findings,
            key=lambda f: (-SEVERITY_ORDER.get(f.severity, 0), f.activity_key, f.code),
        )
        for finding in sorted_findings[:100]:
            lines.append(
                f"| {finding.severity} | {finding.code} | {finding.activity_id or '-'} | "
                f"{finding.message.replace('|', '/')} |"
            )
        if not findings:
            lines.append("| INFO | NONE | - | No findings |")
        lines.extend(
            [
                "",
                "Dashboard RYG is a derived planning aid. RED/YELLOW does not change source workbook authority.",
                "",
            ]
        )
        self.dashboard_md.write_text("\n".join(lines), encoding="utf-8")

        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary["A1"] = "Planning / Schedule Health Dashboard"
        summary["A1"].font = Font(bold=True, size=16)
        summary["A3"] = "As-of"
        summary["B3"] = self.as_of.isoformat()
        summary["A4"] = "Authority"
        summary["B4"] = AUTHORITY
        metrics = [
            ("Activities", kpis["activities_total"]),
            ("GREEN", kpis["ryg"]["green"]),
            ("YELLOW", kpis["ryg"]["yellow"]),
            ("RED", kpis["ryg"]["red"]),
            ("Validation pass rate %", kpis["validation"]["validation_pass_rate_pct"]),
            ("Field completeness %", kpis["planning"]["field_completeness_pct"]),
            ("Overdue", kpis["planning"]["overdue"]),
            ("Negative float", kpis["planning"]["negative_float"]),
            ("Milestones", kpis["planning"]["milestones"]),
            ("Mean progress %", kpis["planning"]["mean_progress_pct"]),
        ]
        summary.append([])
        summary.append(["KPI", "Value"])
        for cell in summary[6]:
            cell.font = Font(bold=True)
        for name, value in metrics:
            summary.append([name, value])
            if name in RYG_FILLS:
                summary.cell(row=summary.max_row, column=1).fill = PatternFill(
                    "solid", fgColor=RYG_FILLS[name]
                )
        summary["A18"] = "Legend"
        summary["A18"].font = Font(bold=True)
        for offset, ryg_name in enumerate(("GREEN", "YELLOW", "RED"), start=19):
            summary.cell(row=offset, column=1).value = ryg_name
            summary.cell(row=offset, column=1).fill = PatternFill(
                "solid", fgColor=RYG_FILLS[ryg_name]
            )
        self._autosize(summary)

        schedule_ws = wb.create_sheet("Schedule")
        schedule_ws.append(CANONICAL_FIELDS)
        for cell in schedule_ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            data = asdict(row)
            data.pop("_predecessor_list", None)
            schedule_ws.append([data.get(field_name) for field_name in CANONICAL_FIELDS])
            schedule_ws.cell(
                row=schedule_ws.max_row,
                column=CANONICAL_FIELDS.index("ryg") + 1,
            ).fill = PatternFill("solid", fgColor=RYG_FILLS[row.ryg])
        schedule_ws.freeze_panes = "A2"
        self._autosize(schedule_ws)

        findings_ws = wb.create_sheet("Findings")
        finding_fields = [
            "source_table_id", "activity_key", "activity_id", "source_row",
            "kind", "severity", "code", "message",
        ]
        findings_ws.append(finding_fields)
        for cell in findings_ws[1]:
            cell.font = Font(bold=True)
        for finding in sorted_findings:
            findings_ws.append([getattr(finding, field_name) for field_name in finding_fields])
        findings_ws.freeze_panes = "A2"
        self._autosize(findings_ws)
        wb.save(self.dashboard_xlsx)

    def _write_manifest(
        self,
        candidate_tables: Sequence[dict[str, Any]],
        rows: Sequence[ScheduleRow],
        findings: Sequence[Finding],
        kpis: dict[str, Any],
    ) -> None:
        outputs = {
            "canonical_csv": self.canonical_csv,
            "canonical_xlsx": self.canonical_xlsx,
            "findings_csv": self.findings_csv,
            "kpis_json": self.kpis_json,
            "dashboard_md": self.dashboard_md,
            "dashboard_xlsx": self.dashboard_xlsx,
        }
        payload = {
            "schema": "gbogeb.schedule_projection_manifest.v1",
            "generated_at": datetime.now().astimezone().isoformat(),
            "as_of": self.as_of.isoformat(),
            "authority": {
                "source_workbook_remains_authoritative_input": True,
                "source_table_manifest_remains_upstream_lineage": True,
                "schedule_projection_is_derived": True,
                "dashboard_is_decision_support": True,
                "authority_transfer": False,
            },
            "source": {
                "table_manifest": str(self.table_manifest_path),
                "table_manifest_sha256": sha256_file(self.table_manifest_path),
                "schedule_candidate_tables": len(candidate_tables),
            },
            "summary": {
                "activities": len(rows),
                "findings": len(findings),
                "red": kpis["ryg"]["red"],
                "yellow": kpis["ryg"]["yellow"],
                "green": kpis["ryg"]["green"],
                "validation_pass_rate_pct": kpis["validation"]["validation_pass_rate_pct"],
                "process_status": "PASS",
            },
            "outputs": {
                key: {
                    "path": str(path.relative_to(self.output_root)),
                    "sha256": sha256_file(path),
                }
                for key, path in outputs.items()
            },
        }
        self.schedule_manifest.write_text(
            json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )

    def run(self) -> Path:
        self._prepare()
        manifest = self._load_manifest()
        candidate_tables = [
            table
            for table in manifest.get("tables", [])
            if table.get("schedule_candidate") is True
        ]
        rows: list[ScheduleRow] = []
        findings: list[Finding] = []

        for table in candidate_tables:
            headers, source_rows = self._read_candidate_rows(table)
            mapping = _map_headers(headers)
            for row_index, values in enumerate(source_rows, start=2):
                if not any(_clean(value) for value in values):
                    continue
                rows.append(
                    self._row_from_source(
                        table,
                        row_index,
                        headers,
                        values,
                        mapping,
                        findings,
                    )
                )

        self._validate_ids_and_dependencies(rows, findings)
        self._apply_risk_and_ryg(rows, findings)
        kpis = self._calculate_kpis(rows, findings, len(candidate_tables))
        self._write_csvs(rows, findings)
        self._write_canonical_xlsx(rows, findings)
        self.kpis_json.write_text(
            json.dumps(kpis, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        self._write_dashboard(rows, findings, kpis)
        self._write_manifest(candidate_tables, rows, findings, kpis)
        return self.schedule_manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Build canonical schedule projection, validation and RYG dashboard."
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=Path("."),
        help="Pipeline root containing Outputs/excel/table_manifest.json",
    )
    parser.add_argument(
        "--table-manifest",
        type=Path,
        default=None,
        help="Optional explicit table_manifest.json path",
    )
    parser.add_argument(
        "--as-of",
        type=date.fromisoformat,
        default=None,
        help="Deterministic planning as-of date (YYYY-MM-DD)",
    )
    parser.add_argument(
        "--due-soon-days",
        type=int,
        default=14,
        help="Yellow due-soon horizon in calendar days",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    engine = ScheduleProjectionEngine(
        output_root=args.output_root,
        table_manifest_path=args.table_manifest,
        as_of=args.as_of,
        due_soon_days=args.due_soon_days,
    )
    manifest_path = engine.run()
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    print(
        json.dumps(
            {
                "status": payload["summary"]["process_status"],
                "activities": payload["summary"]["activities"],
                "red": payload["summary"]["red"],
                "yellow": payload["summary"]["yellow"],
                "green": payload["summary"]["green"],
                "manifest": str(manifest_path),
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
